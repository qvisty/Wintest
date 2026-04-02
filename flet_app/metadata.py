"""Metadata extraction for common file types."""

import os
from datetime import datetime


def extract_metadata(path: str) -> dict:
    """Extract metadata from a file. Returns a dict of key-value pairs."""
    ext = os.path.splitext(path)[1].lower()

    meta = _basic_metadata(path)

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".doc": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".xls": _extract_xlsx,
        ".pptx": _extract_pptx,
        ".msg": _extract_msg,
        ".jpg": _extract_image,
        ".jpeg": _extract_image,
        ".png": _extract_image,
        ".gif": _extract_image,
        ".bmp": _extract_image,
        ".tiff": _extract_image,
        ".tif": _extract_image,
    }

    extractor = extractors.get(ext)
    if extractor:
        try:
            meta.update(extractor(path))
        except Exception as e:
            meta["Metadata-fejl"] = str(e)

    return meta


def _basic_metadata(path: str) -> dict:
    """Basic file system metadata available for all files."""
    try:
        stat = os.stat(path)
        return {
            "Filstørrelse": _fmt_size(stat.st_size),
            "Oprettet": datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M"),
            "Ændret": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "Sidst åbnet": datetime.fromtimestamp(stat.st_atime).strftime("%Y-%m-%d %H:%M"),
        }
    except OSError:
        return {}


def _fmt_size(size_bytes: int) -> str:
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    if size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
    return f"{size_bytes / (1024 * 1024 * 1024):.1f} GB"


def _extract_pdf(path: str) -> dict:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        return {"Note": "Installér PyPDF2 for PDF-metadata"}

    reader = PdfReader(path)
    meta = {}
    meta["Sider"] = str(len(reader.pages))

    info = reader.metadata
    if info:
        if info.title:
            meta["Titel"] = info.title
        if info.author:
            meta["Forfatter"] = info.author
        if info.subject:
            meta["Emne"] = info.subject
        if info.creator:
            meta["Oprettet med"] = info.creator
        if info.creation_date:
            meta["PDF oprettet"] = str(info.creation_date)

    return meta


def _extract_docx(path: str) -> dict:
    try:
        from docx import Document
    except ImportError:
        return {"Note": "Installér python-docx for Word-metadata"}

    doc = Document(path)
    props = doc.core_properties
    meta = {}

    if props.title:
        meta["Titel"] = props.title
    if props.author:
        meta["Forfatter"] = props.author
    if props.subject:
        meta["Emne"] = props.subject
    if props.created:
        meta["Oprettet"] = str(props.created)
    if props.modified:
        meta["Ændret"] = str(props.modified)
    if props.last_modified_by:
        meta["Sidst ændret af"] = props.last_modified_by
    if props.revision:
        meta["Revision"] = str(props.revision)
    if props.category:
        meta["Kategori"] = props.category
    if props.comments:
        meta["Kommentarer"] = props.comments

    meta["Afsnit"] = str(len(doc.paragraphs))
    meta["Tabeller"] = str(len(doc.tables))

    return meta


def _extract_xlsx(path: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"Note": "Installér openpyxl for Excel-metadata"}

    wb = load_workbook(path, read_only=True, data_only=True)
    props = wb.properties
    meta = {}

    if props.title:
        meta["Titel"] = props.title
    if props.creator:
        meta["Forfatter"] = props.creator
    if props.subject:
        meta["Emne"] = props.subject
    if props.created:
        meta["Oprettet"] = str(props.created)
    if props.modified:
        meta["Ændret"] = str(props.modified)
    if props.lastModifiedBy:
        meta["Sidst ændret af"] = props.lastModifiedBy

    meta["Ark"] = ", ".join(wb.sheetnames)
    meta["Antal ark"] = str(len(wb.sheetnames))

    wb.close()
    return meta


def _extract_pptx(path: str) -> dict:
    try:
        from pptx import Presentation
    except ImportError:
        return {"Note": "Installér python-pptx for PowerPoint-metadata"}

    prs = Presentation(path)
    props = prs.core_properties
    meta = {}

    if props.title:
        meta["Titel"] = props.title
    if props.author:
        meta["Forfatter"] = props.author
    if props.subject:
        meta["Emne"] = props.subject
    if props.created:
        meta["Oprettet"] = str(props.created)
    if props.modified:
        meta["Ændret"] = str(props.modified)

    meta["Slides"] = str(len(prs.slides))

    return meta


def _extract_msg(path: str) -> dict:
    try:
        import extract_msg
    except ImportError:
        return {"Note": "Installér extract-msg for email-metadata"}

    msg = extract_msg.Message(path)
    meta = {}

    if msg.subject:
        meta["Emne"] = msg.subject
    if msg.sender:
        meta["Afsender"] = msg.sender
    if msg.to:
        meta["Til"] = msg.to
    if msg.date:
        meta["Dato"] = str(msg.date)
    if msg.cc:
        meta["CC"] = msg.cc

    attachments = msg.attachments
    if attachments:
        meta["Vedhæftninger"] = str(len(attachments))
        meta["Vedhæftede filer"] = ", ".join(a.longFilename or a.shortFilename or "?" for a in attachments)

    msg.close()
    return meta


def _extract_image(path: str) -> dict:
    try:
        from PIL import Image
        from PIL.ExifTags import TAGS
    except ImportError:
        return {"Note": "Installér Pillow for billed-metadata"}

    img = Image.open(path)
    meta = {
        "Dimensioner": f"{img.width} x {img.height} px",
        "Format": img.format or "Ukendt",
        "Farvetilstand": img.mode,
    }

    exif_data = img.getexif()
    if exif_data:
        interesting_tags = {
            "Make": "Kamera",
            "Model": "Kameramodel",
            "DateTime": "Foto taget",
            "DateTimeOriginal": "Foto taget",
            "Software": "Software",
            "ImageDescription": "Beskrivelse",
            "Artist": "Fotograf",
            "Copyright": "Copyright",
        }
        for tag_id, value in exif_data.items():
            tag_name = TAGS.get(tag_id, str(tag_id))
            if tag_name in interesting_tags:
                meta[interesting_tags[tag_name]] = str(value)

    img.close()
    return meta
